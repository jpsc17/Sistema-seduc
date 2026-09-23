# -*- coding: utf-8 -*-
r"""
================================================================================
SISTEMA SEDUC - PIPELINE DE EXTRAÇÃO DE RESULTADOS EM PDF
================================================================================
Este script realiza:
1. Leitura e extração estruturada dos 9 relatórios em PDF da SEDUC.
2. Geração da planilha Excel consolidada 'Seduc_Resultados_Consolidado.xlsx'
   com abas individuais para cada PDF, formatação visual profissional,
   filtros automáticos e aba de Visão Geral.
3. Preparação completa de dados para PostgreSQL:
   - CSVs limpos em UTF-8 (sem problemas de delimitador, NULLs padronizados).
   - Script DDL 'schema_seduc.sql' com CREATE TABLE, tipos adequados e índices.
   - Script de carga 'load_data.sql' (via comando \copy do PostgreSQL).
   - Script 'inserts_seduc.sql' com comandos INSERT prontos para execução em
     qualquer cliente (pgAdmin, DBeaver, psql, Supabase, RDS, etc.).
================================================================================
"""

import os
import re
import glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import pdfplumber

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PDF_DIR = os.path.join(BASE_DIR, "resultadospdf")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "Seduc_Resultados_Consolidado.xlsx")
PG_DIR = os.path.join(BASE_DIR, "postgres")
PG_DATA_DIR = os.path.join(PG_DIR, "data")


# ----------------------------------------------------------------------
# FUNÇÕES DE LIMPEZA E PADRONIZAÇÃO DE DADOS
# ----------------------------------------------------------------------

def clean_text(v):
    if v is None:
        return None
    s = str(v).replace('\r', ' ').replace('\n', ' ').strip()
    s = re.sub(r'\s+', ' ', s)
    return s if s else None

def parse_decimal(v):
    if v is None:
        return None
    s = clean_text(v)
    if not s or s in ['-', '--', '.', 'None', 'null']:
        return None
    # Trata espacos ao redor da virgula como '0 ,6' -> '0.6'
    s = re.sub(r'\s*,\s*', '.', s)
    s = s.replace(' ', '')
    try:
        return round(float(s), 4)
    except ValueError:
        return None

def parse_int(v):
    if v is None:
        return None
    s = clean_text(v)
    if not s or s in ['-', '--', '.', 'None', 'null']:
        return None
    # Remove pontos de milhar e espaços
    s = s.replace('.', '').replace(' ', '')
    try:
        return int(s)
    except ValueError:
        return None

def parse_bool(v):
    if v is None:
        return None
    s = clean_text(v)
    if not s:
        return None
    s_upper = s.upper()
    if 'SIM' in s_upper:
        return True
    if 'NÃO' in s_upper or 'NAO' in s_upper:
        return False
    return s


# ----------------------------------------------------------------------
# EXTRATORES PARA CADA UM DOS 9 PDFs
# ----------------------------------------------------------------------

def extract_eja_aee(fpath):
    """1. Escolas de EJA e AEE - Bônus Professor e Administrativo (Escola Exclusiva)"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    if r[1] and 'Munic' in str(r[1]):
                        continue
                    if r[7] and ('JOVENS' in str(r[7]) or 'EF_ANOS' in str(r[7])):
                        continue
                    code = clean_text(r[4])
                    if code and code.isdigit():
                        rows.append({
                            'regional': clean_text(r[0]),
                            'municipio': clean_text(r[1]),
                            'localizacao': clean_text(r[2]),
                            'escola_indigena': parse_bool(r[3]),
                            'codigo_escola': code,
                            'tem_publicacao': parse_bool(r[5]),
                            'nome_escola': clean_text(r[6]),
                            'eja_fundamental_iniciais': parse_decimal(r[7]),
                            'eja_fundamental_finais': parse_decimal(r[8]),
                            'eja_medio': parse_decimal(r[9]),
                            'atendimento_especializado_aee': parse_decimal(r[10])
                        })
    return pd.DataFrame(rows)


def extract_nao_publ_regular_admin(fpath):
    """2. Escolas Não Publicadas - Ensino Regular - Cargo Administrativo"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    if r[1] and 'Munic' in str(r[1]):
                        continue
                    code = clean_text(r[4])
                    if code and code.isdigit():
                        rows.append({
                            'regional': clean_text(r[0]),
                            'municipio': clean_text(r[1]),
                            'localizacao': clean_text(r[2]),
                            'escola_indigena': parse_bool(r[3]),
                            'codigo_escola': code,
                            'nome_escola': clean_text(r[5]),
                            'indice_bonus_admin': parse_decimal(r[6])
                        })
    return pd.DataFrame(rows)


def extract_nao_publ_regular_prof(fpath):
    """3. Escolas Não Publicadas - Ensino Regular - Bônus Professores em Sala de Aula"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            txt = p.extract_text() or ''
            if 'ANOS INICIAIS' in txt.upper():
                etapa = 'ENSINO FUNDAMENTAL ANOS INICIAIS'
            elif 'ANOS FINAIS' in txt.upper():
                etapa = 'ENSINO FUNDAMENTAL ANOS FINAIS'
            elif 'MÉDIO' in txt.upper() or 'MEDIO' in txt.upper():
                etapa = 'ENSINO MEDIO'
            else:
                etapa = 'ENSINO REGULAR'
            
            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    if r[1] and 'Munic' in str(r[1]):
                        continue
                    if r[3] and 'Ind' in str(r[3]):
                        continue
                    code = clean_text(r[4])
                    if code and code.isdigit():
                        rows.append({
                            'regional': clean_text(r[0]),
                            'municipio': clean_text(r[1]),
                            'localizacao': clean_text(r[2]),
                            'escola_indigena': parse_bool(r[3]),
                            'codigo_escola': code,
                            'nome_escola': clean_text(r[5]),
                            'ponto_bonus_professor': parse_decimal(r[6]),
                            'etapa_ensino': etapa
                        })
    return pd.DataFrame(rows)


def extract_nao_publ_sectet(fpath):
    """4. Escolas Não Publicadas SECTET - Bônus Professor e Cargo Administrativo"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    if r[1] and 'Munic' in str(r[1]):
                        continue
                    code = clean_text(r[4])
                    if code and code.isdigit():
                        rows.append({
                            'regional': clean_text(r[0]),
                            'municipio': clean_text(r[1]),
                            'localizacao': clean_text(r[2]),
                            'escola_indigena': parse_bool(r[3]),
                            'codigo_escola': code,
                            'nome_escola': clean_text(r[5]),
                            'indice_bonus_gestao': parse_decimal(r[6])
                        })
    return pd.DataFrame(rows)


def extract_publ_regular_admin(fpath):
    """5. Escolas Publicadas - Bônus Cargos Administrativos"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    if r[0] and ('Munic' in str(r[0]) or 'MATR' in str(r[0])):
                        continue
                    code = clean_text(r[2])
                    if code and code.isdigit():
                        rows.append({
                            'municipio': clean_text(r[0]),
                            'localizacao': clean_text(r[1]),
                            'codigo_escola': code,
                            'nome_escola': clean_text(r[3]),
                            'matricula_ef_iniciais': parse_int(r[4]),
                            'matricula_ef_finais': parse_int(r[5]),
                            'matricula_em': parse_int(r[6]),
                            'matricula_total': parse_int(r[7]),
                            'pontos_ef_iniciais': parse_decimal(r[8]),
                            'pontos_ef_finais': parse_decimal(r[9]),
                            'pontos_em': parse_decimal(r[10]),
                            'ponderado_ef_iniciais': parse_decimal(r[11]),
                            'ponderado_ef_finais': parse_decimal(r[12]),
                            'ponderado_em': parse_decimal(r[13]),
                            'indice_bonus': parse_decimal(r[14])
                        })
    return pd.DataFrame(rows)


def extract_publ_regular_prof(fpath):
    """6. Escolas Publicadas - Bônus Professor em Sala de Aula"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            txt = p.extract_text() or ''
            first_lines = txt.split('\n')[:3]
            txt_header = ' '.join(first_lines).upper()
            if 'ANOS INICIAIS' in txt_header:
                etapa = 'ENSINO FUNDAMENTAL ANOS INICIAIS'
            elif 'ANOS FINAIS' in txt_header:
                etapa = 'ENSINO FUNDAMENTAL ANOS FINAIS'
            elif 'MÉDIO' in txt_header or 'MEDIO' in txt_header:
                etapa = 'ENSINO MEDIO'
            else:
                etapa = 'ENSINO REGULAR'

            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    code = clean_text(r[2]) if len(r) > 2 else None
                    if code and code.isdigit():
                        if etapa == 'ENSINO FUNDAMENTAL ANOS INICIAIS':
                            rows.append({
                                'municipio': clean_text(r[0]),
                                'regiao_integracao': clean_text(r[1]),
                                'codigo_escola': code,
                                'nome_escola': clean_text(r[3]),
                                'etapa_ensino': etapa,
                                'meta_pactuada': parse_decimal(r[4]),
                                'decimo_crescimento': parse_decimal(r[5]),
                                'ponto_crescimento': parse_decimal(r[6]),
                                'ponto_regiao_integracao': parse_decimal(r[7]),
                                'ponto_fluxo': parse_decimal(r[8]),
                                'ponto_alfabetizacao': parse_decimal(r[9]),
                                'bonus_prof_1_e_2_ano': parse_decimal(r[10]),
                                'bonus_prof_3_a_5_ano': parse_decimal(r[11]),
                                'indice_bonus': None
                            })
                        else:
                            rows.append({
                                'municipio': clean_text(r[0]),
                                'regiao_integracao': clean_text(r[1]),
                                'codigo_escola': code,
                                'nome_escola': clean_text(r[3]),
                                'etapa_ensino': etapa,
                                'meta_pactuada': parse_decimal(r[4]),
                                'decimo_crescimento': parse_decimal(r[5]),
                                'ponto_crescimento': parse_decimal(r[6]),
                                'ponto_regiao_integracao': parse_decimal(r[7]),
                                'ponto_fluxo': parse_decimal(r[8]),
                                'ponto_alfabetizacao': None,
                                'bonus_prof_1_e_2_ano': None,
                                'bonus_prof_3_a_5_ano': None,
                                'indice_bonus': parse_decimal(r[9]) if len(r) > 9 else None
                            })
    return pd.DataFrame(rows)


def extract_publ_sectet(fpath):
    """7. Escolas Publicadas SECTET - Bônus Professor e Cargo Administrativo"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    code = clean_text(r[1]) if len(r) > 1 else None
                    if code and code.isdigit():
                        rows.append({
                            'municipio': clean_text(r[0]),
                            'codigo_escola': code,
                            'sectet': parse_bool(r[2]),
                            'nome_escola': clean_text(r[3]),
                            'atingiu_meta_pactuada': parse_decimal(r[4]),
                            'decimo_crescimento': parse_decimal(r[5]),
                            'ponto_crescimento': parse_decimal(r[6]),
                            'fluxo': parse_decimal(r[7]),
                            'bonus_prof_vinculado_turma': parse_decimal(r[8]),
                            'bonus_cargo_administrativo': parse_decimal(r[9])
                        })
    return pd.DataFrame(rows)


def extract_pontos_bonus_dre(fpath):
    """8. Pontos de Bônus por Diretoria Regional de Ensino (DRE)"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    dre = clean_text(r[1]) if len(r) > 1 else None
                    if not dre or dre in ['DRE', '', 'None']:
                        continue
                    rows.append({
                        'ordem': len(rows) + 1,
                        'dre': dre,
                        'pontos_ai': parse_decimal(r[2]),
                        'pontos_af': parse_decimal(r[3]),
                        'pontos_em': parse_decimal(r[4]),
                        'matricula_ai': parse_int(r[5]),
                        'matricula_af': parse_int(r[6]),
                        'matricula_em': parse_int(r[7]),
                        'matricula_total': parse_int(r[8]),
                        'bonus_ai': parse_decimal(r[9]),
                        'bonus_af': parse_decimal(r[10]),
                        'bonus_em': parse_decimal(r[11]),
                        'bonus_ri': parse_decimal(r[12]),
                        'bonus_total': parse_decimal(r[13])
                    })
    return pd.DataFrame(rows)


def extract_ideb_dre(fpath):
    """9. Índice de Desenvolvimento da Educação Básica (IDEB) por DRE"""
    rows = []
    with pdfplumber.open(fpath) as pdf:
        for p in pdf.pages:
            txt = p.extract_text() or ''
            if 'ANOS INICIAIS' in txt.upper():
                etapa = 'ENSINO FUNDAMENTAL ANOS INICIAIS'
            elif 'ANOS FINAIS' in txt.upper():
                etapa = 'ENSINO FUNDAMENTAL ANOS FINAIS'
            elif 'MÉDIO' in txt.upper() or 'MEDIO' in txt.upper():
                etapa = 'ENSINO MEDIO'
            else:
                etapa = 'ENSINO GERAL'

            tables = p.extract_tables()
            for t in tables:
                for r in t:
                    dre = clean_text(r[1]) if len(r) > 1 else None
                    if not dre or dre in ['DRE', '', 'None']:
                        continue
                    rows.append({
                        'ordem': len(rows) + 1,
                        'dre': dre,
                        'etapa_ensino': etapa,
                        'desempenho_lingua_portuguesa': parse_decimal(r[2]),
                        'desempenho_matematica': parse_decimal(r[3]),
                        'nota_padronizada_lp': parse_decimal(r[4]),
                        'nota_padronizada_mat': parse_decimal(r[5]),
                        'nota_padronizada_media': parse_decimal(r[6]),
                        'fluxo_tempo_medio': parse_decimal(r[7]),
                        'ideb': parse_decimal(r[8])
                    })
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------
# MAPEAMENTO DOS 9 ARQUIVOS
# ----------------------------------------------------------------------

DATASET_CONFIG = [
    {
        "filename": "ESCOLAS DE EDUCAÇAO DE JOVENS E ADULTOS E AEE - BONUS PARA PROFESSOR E ADMINISTRATIVO_ESCOLA_EXCLUSIVA.pdf",
        "sheet_name": "EJA_AEE_Exclusiva",
        "table_name": "seduc_bonus_eja_aee",
        "title": "Bônus EJA e AEE - Professor e Administrativo (Escola Exclusiva)",
        "extractor": extract_eja_aee
    },
    {
        "filename": "ESCOLAS NÃO PUBLICADAS_ENSINO REGULAR _CARGO ADMINISTRATIVO.pdf",
        "sheet_name": "NaoPubl_Reg_Admin",
        "table_name": "seduc_nao_publ_regular_admin",
        "title": "Escolas Não Publicadas - Ensino Regular - Cargo Administrativo",
        "extractor": extract_nao_publ_regular_admin
    },
    {
        "filename": "ESCOLAS NÃO PUBLICADAS_ENSINO REGULAR_BÔNUS PROFESSORES EM SALA DE AULA.pdf",
        "sheet_name": "NaoPubl_Reg_Prof",
        "table_name": "seduc_nao_publ_regular_prof",
        "title": "Escolas Não Publicadas - Ensino Regular - Bônus Professores em Sala",
        "extractor": extract_nao_publ_regular_prof
    },
    {
        "filename": "ESCOLAS NÃO PUBLICADAS_SECTET - BÔNUS PROFESSOR E CARGO ADMINISTRATIVO.pdf",
        "sheet_name": "NaoPubl_SECTET",
        "table_name": "seduc_nao_publ_sectet",
        "title": "Escolas Não Publicadas SECTET - Bônus Professor e Cargo Administrativo",
        "extractor": extract_nao_publ_sectet
    },
    {
        "filename": "ESCOLAS PUBLICADAS - BÔNUS CARGOS ADMINISTRATIVOS.pdf",
        "sheet_name": "Publ_Reg_Admin",
        "table_name": "seduc_publ_regular_admin",
        "title": "Escolas Publicadas - Ensino Regular - Bônus Cargos Administrativos",
        "extractor": extract_publ_regular_admin
    },
    {
        "filename": "ESCOLAS PUBLICADAS_BÔNUS PROFESSOR EM SALA DE AULA.pdf",
        "sheet_name": "Publ_Reg_Prof",
        "table_name": "seduc_publ_regular_prof",
        "title": "Escolas Publicadas - Bônus Professor em Sala de Aula (Iniciais, Finais, Médio)",
        "extractor": extract_publ_regular_prof
    },
    {
        "filename": "ESCOLAS PUBLICADAS_SECTET_BÔNUS PROFESSOR EM SALA DE AULA E CARGO ADMINISTRATIVO.pdf",
        "sheet_name": "Publ_SECTET",
        "table_name": "seduc_publ_sectet",
        "title": "Escolas Publicadas SECTET - Bônus Professor e Cargo Administrativo",
        "extractor": extract_publ_sectet
    },
    {
        "filename": "PONTOS DE BONOS_DRE.pdf",
        "sheet_name": "Pontos_Bonus_DRE",
        "table_name": "seduc_pontos_bonus_dre",
        "title": "Pontos de Bônus por Diretoria Regional de Ensino (DRE)",
        "extractor": extract_pontos_bonus_dre
    },
    {
        "filename": "Índice de Desenvolvimento da Educação Básica_DRE.pdf",
        "sheet_name": "IDEB_DRE",
        "table_name": "seduc_ideb_dre",
        "title": "Índice de Desenvolvimento da Educação Básica (IDEB) por DRE",
        "extractor": extract_ideb_dre
    }
]


# ----------------------------------------------------------------------
# ESTILIZAÇÃO DO EXCEL
# ----------------------------------------------------------------------

COLOR_HEADER_BG = "1F4E78"    # Azul Marinho Profissional
COLOR_HEADER_TXT = "FFFFFF"   # Branco
COLOR_ZEBRA_BG = "F2F5F9"     # Cinza-Azulado Claro
COLOR_BORDER = "D9D9D9"       # Cinza Suave

font_header = Font(name="Segoe UI", size=11, bold=True, color=COLOR_HEADER_TXT)
font_title = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
font_data = Font(name="Segoe UI", size=10)
font_bold = Font(name="Segoe UI", size=10, bold=True)

fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
fill_zebra = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ----------------------------------------------------------------------
# MAPEAMENTO DE NOMES DE COLUNAS PARA EXIBIÇÃO NO EXCEL
# ----------------------------------------------------------------------

COLUMN_DISPLAY_NAMES = {
    'regional': 'Regional',
    'municipio': 'Nome do Município',
    'localizacao': 'Localização',
    'escola_indigena': 'Escola Indígena',
    'codigo_escola': 'Código da Escola',
    'tem_publicacao': 'Tem Publicação',
    'nome_escola': 'Nome da Escola',
    'eja_fundamental_iniciais': 'EJA - EF Anos Iniciais',
    'eja_fundamental_finais': 'EJA - EF Anos Finais',
    'eja_medio': 'EJA - Ensino Médio',
    'atendimento_especializado_aee': 'Atendimento Especializado (AEE)',
    'indice_bonus_admin': 'Índice de Bônus Cargo Administrativo',
    'ponto_bonus_professor': 'Ponto de Bônus Professor',
    'etapa_ensino': 'Etapa de Ensino',
    'indice_bonus_gestao': 'Índice de Bônus Gestão (Admin e Professores)',
    'matricula_ef_iniciais': 'Matrícula EF Anos Iniciais',
    'matricula_ef_finais': 'Matrícula EF Anos Finais',
    'matricula_em': 'Matrícula Ensino Médio',
    'matricula_total': 'Matrícula Total',
    'pontos_ef_iniciais': 'Pontos Bônus EF Iniciais',
    'pontos_ef_finais': 'Pontos Bônus EF Finais',
    'pontos_em': 'Pontos Bônus Ensino Médio',
    'ponderado_ef_iniciais': 'Pontos Ponderados EF Iniciais',
    'ponderado_ef_finais': 'Pontos Ponderados EF Finais',
    'ponderado_em': 'Pontos Ponderados Ensino Médio',
    'indice_bonus': 'Índice de Bônus Final',
    'regiao_integracao': 'Região de Integração',
    'meta_pactuada': 'Atingiu Meta Pactuada',
    'atingiu_meta_pactuada': 'Atingiu Meta Pactuada',
    'decimo_crescimento': 'Décimo Crescimento',
    'ponto_crescimento': 'Ponto Crescimento',
    'ponto_regiao_integracao': 'Ponto Região Integração',
    'ponto_fluxo': 'Ponto Fluxo',
    'ponto_alfabetizacao': 'Ponto Alfabetização',
    'bonus_prof_1_e_2_ano': 'Bônus Professor (1º e 2º ano)',
    'bonus_prof_3_a_5_ano': 'Bônus Professor (3º a 5º ano)',
    'bonus_prof_vinculado_turma': 'Bônus Professor em Turma',
    'bonus_cargo_administrativo': 'Bônus Cargo Administrativo',
    'sectet': 'SECTET',
    'fluxo': 'Fluxo',
    'ordem': 'Ordem',
    'dre': 'DRE (Diretoria Regional)',
    'pontos_ai': 'Pontos EF Anos Iniciais',
    'pontos_af': 'Pontos EF Anos Finais',
    'pontos_em': 'Pontos Ensino Médio',
    'matricula_ai': 'Matrícula EF Anos Iniciais',
    'matricula_af': 'Matrícula EF Anos Finais',
    'bonus_ai': 'Bônus EF Anos Iniciais',
    'bonus_af': 'Bônus EF Anos Finais',
    'bonus_ri': 'Bônus Região Integração',
    'bonus_total': 'Bônus Total DRE',
    'desempenho_lingua_portuguesa': 'Desempenho Língua Portuguesa',
    'desempenho_matematica': 'Desempenho Matemática',
    'nota_padronizada_lp': 'Nota Padronizada LP',
    'nota_padronizada_mat': 'Nota Padronizada Matemática',
    'nota_padronizada_media': 'Média Nota Padronizada',
    'fluxo_tempo_medio': 'Fluxo (Tempo Médio)',
    'ideb': 'IDEB'
}

def get_col_display(cname):
    return COLUMN_DISPLAY_NAMES.get(cname, cname.replace('_', ' ').title())


def style_worksheet(ws, df, title_text):
    # A Linha 1 é o cabeçalho direto com nomes claros e visíveis
    header_row_idx = 1
    ws.row_dimensions[header_row_idx].height = 32

    # Formata cabeçalho na Linha 1 com nomes em português
    for col_idx, col_name in enumerate(df.columns, 1):
        display_name = get_col_display(col_name)
        cell = ws.cell(row=header_row_idx, column=col_idx, value=display_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = thin_border

    # Congela o cabeçalho na Linha 1 para que NUNCA desapareça ao rolar
    ws.freeze_panes = "A2"

    # Formata linhas de dados a partir da Linha 2
    start_data_row = 2
    end_data_row = start_data_row + len(df) - 1

    for row_idx in range(start_data_row, end_data_row + 1):
        ws.row_dimensions[row_idx].height = 20
        is_zebra = (row_idx % 2 == 0)
        current_fill = fill_zebra if is_zebra else fill_white
        
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_data
            cell.fill = current_fill
            cell.border = thin_border
            
            val = cell.value
            col_name = df.columns[col_idx - 1]
            
            # Formatações específicas por tipo de coluna
            if isinstance(val, (int, float)):
                if 'matricula' in col_name or col_name in ['ordem']:
                    cell.number_format = '#,##0'
                    cell.alignment = align_right
                else:
                    cell.number_format = '0.00'
                    cell.alignment = align_right
            elif isinstance(val, bool):
                cell.value = "SIM" if val else "NÃO"
                cell.alignment = align_center
            elif col_name == 'codigo_escola':
                cell.number_format = '@'
                cell.alignment = align_center
            elif col_name in ['escola_indigena', 'tem_publicacao', 'sectet']:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Ativa autofiltro na linha 1
    if len(df) > 0:
        last_col_letter = get_column_letter(len(df.columns))
        ws.auto_filter.ref = f"A1:{last_col_letter}{end_data_row}"

    # Auto-ajuste da largura das colunas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


def create_overview_sheet(wb, summary_data):
    """Cria a aba inicial '00_Visao_Geral' com sumário de todos os relatórios"""
    ws = wb.create_sheet(title="00_Visao_Geral", index=0)
    ws.views.sheetView[0].showGridLines = True
    
    ws.cell(row=1, column=1, value="SISTEMA SEDUC - PAINEL CONSOLIDADO DE EXTRAÇÃO").font = font_title
    ws.cell(row=2, column=1, value="Relatórios de Índices de Bônus e IDEB - Extraídos dos Documentos Oficiais em PDF").font = Font(name="Segoe UI", size=10, italic=True, color="555555")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 12

    headers = ["#", "Nome da Aba", "Tabela PostgreSQL", "Descrição do Relatório", "Total de Registros", "Arquivo PDF de Origem"]
    header_row = 4
    ws.row_dimensions[header_row].height = 28
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = thin_border

    total_records = 0
    for idx, item in enumerate(summary_data, 1):
        r_idx = header_row + idx
        ws.row_dimensions[r_idx].height = 22
        current_fill = fill_zebra if idx % 2 == 0 else fill_white
        
        row_vals = [
            idx,
            item["sheet_name"],
            item["table_name"],
            item["title"],
            item["records"],
            item["filename"]
        ]
        total_records += item["records"]
        
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.fill = current_fill
            cell.border = thin_border
            if col_idx in [1, 5]:
                cell.alignment = align_right
                if col_idx == 5:
                    cell.number_format = '#,##0'
            elif col_idx in [2, 3]:
                cell.alignment = align_left
                cell.font = font_bold
            else:
                cell.alignment = align_left

    # Linha Totalizadora
    tot_row = header_row + len(summary_data) + 1
    ws.row_dimensions[tot_row].height = 24
    ws.cell(row=tot_row, column=3, value="TOTAL GERAL DE REGISTROS").font = font_bold
    ws.cell(row=tot_row, column=3).alignment = align_right
    tot_cell = ws.cell(row=tot_row, column=5, value=total_records)
    tot_cell.font = font_bold
    tot_cell.alignment = align_right
    tot_cell.number_format = '#,##0'

    # Auto largura
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2, 3]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


# ----------------------------------------------------------------------
# GERAÇÃO DOS ARTEFATOS POSTGRESQL (DDL, CSVs, INSERTS)
# ----------------------------------------------------------------------

def generate_postgres_artifacts(extracted_datasets):
    os.makedirs(PG_DATA_DIR, exist_ok=True)
    
    schema_sql_path = os.path.join(PG_DIR, "schema_seduc.sql")
    load_sql_path = os.path.join(PG_DIR, "load_data.sql")
    inserts_sql_path = os.path.join(PG_DIR, "inserts_seduc.sql")

    schema_statements = []
    load_statements = []
    
    schema_statements.append("""-- ====================================================================
-- SISTEMA SEDUC - ESQUEMA DE BANCO DE DADOS POSTGRESQL
-- Gerado automaticamente a partir dos relatórios em PDF
-- ====================================================================

CREATE SCHEMA IF NOT EXISTS seduc;
SET search_path TO seduc, public;
""")

    load_statements.append("""-- ====================================================================
-- SCRIPT DE CARGA VIA \\copy (PostgreSQL psql)
-- Execute no terminal: psql -U seu_usuario -d seu_banco -f load_data.sql
-- ====================================================================

SET search_path TO seduc, public;
""")

    with open(inserts_sql_path, "w", encoding="utf-8") as f_ins:
        f_ins.write("""-- ====================================================================
-- SISTEMA SEDUC - CARGA DE DADOS VIA INSERT INTO (Compatível com pgAdmin/DBeaver)
-- ====================================================================

CREATE SCHEMA IF NOT EXISTS seduc;
SET search_path TO seduc, public;

BEGIN;
""")

        for item in extracted_datasets:
            table_name = item["table_name"]
            df = item["df"]
            csv_filename = f"{table_name}.csv"
            csv_path = os.path.join(PG_DATA_DIR, csv_filename)

            # Exporta CSV para pasta postgres/data/
            # NULLs são exportados como string vazia (padrão CSV do postgres)
            df.to_csv(csv_path, index=False, encoding="utf-8")

            # Mapeamento dinâmico de tipos PostgreSQL baseado nas colunas do DataFrame
            col_definitions = ["id BIGSERIAL PRIMARY KEY"]
            columns_list = list(df.columns)
            
            for col in columns_list:
                sample_series = df[col].dropna()
                if 'matricula' in col or col in ['ordem']:
                    pg_type = "INTEGER"
                elif 'codigo_escola' in col:
                    pg_type = "VARCHAR(20)"
                elif col in ['escola_indigena', 'tem_publicacao', 'sectet']:
                    pg_type = "BOOLEAN"
                elif any(k in col for k in ['bonus', 'ponto', 'indice', 'ponderado', 'desempenho', 'nota', 'fluxo', 'meta', 'crescimento', 'eja', 'atendimento']):
                    pg_type = "NUMERIC(8,4)"
                elif 'nome' in col or 'municipio' in col or 'regional' in col or 'dre' in col or 'etapa' in col:
                    pg_type = "VARCHAR(255)"
                else:
                    pg_type = "VARCHAR(100)"
                
                col_definitions.append(f"    {col} {pg_type}")

            # DDL da tabela
            cols_ddl = ",\n".join(col_definitions)
            ddl = f"""DROP TABLE IF EXISTS seduc.{table_name} CASCADE;
CREATE TABLE seduc.{table_name} (
{cols_ddl},
    criado_em TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP
);

-- Índices recomendados
"""
            if 'codigo_escola' in columns_list:
                ddl += f"CREATE INDEX idx_{table_name}_codigo_escola ON seduc.{table_name}(codigo_escola);\n"
            if 'municipio' in columns_list:
                ddl += f"CREATE INDEX idx_{table_name}_municipio ON seduc.{table_name}(municipio);\n"
            if 'regional' in columns_list:
                ddl += f"CREATE INDEX idx_{table_name}_regional ON seduc.{table_name}(regional);\n"
            if 'dre' in columns_list:
                ddl += f"CREATE INDEX idx_{table_name}_dre ON seduc.{table_name}(dre);\n"
            if 'etapa_ensino' in columns_list:
                ddl += f"CREATE INDEX idx_{table_name}_etapa ON seduc.{table_name}(etapa_ensino);\n"

            schema_statements.append(ddl)

            # Comando \copy
            cols_joined = ", ".join(columns_list)
            # Use forward slashes in copy path for psql cross-platform compatibility
            csv_path_sql = csv_path.replace("\\", "/")
            load_statements.append(f"\\copy seduc.{table_name} ({cols_joined}) FROM '{csv_path_sql}' WITH (FORMAT csv, HEADER true, NULL '');\n")

            # Escreve INSERTS no arquivo inserts_seduc.sql em blocos
            f_ins.write(f"\n-- Carga da tabela seduc.{table_name} ({len(df)} registros)\n")
            if len(df) > 0:
                batch_size = 100
                cols_sql = ", ".join(columns_list)
                for start_idx in range(0, len(df), batch_size):
                    batch = df.iloc[start_idx:start_idx+batch_size]
                    values_rows = []
                    for _, row in batch.iterrows():
                        row_vals = []
                        for col in columns_list:
                            val = row[col]
                            if pd.isna(val) or val is None:
                                row_vals.append("NULL")
                            elif isinstance(val, bool):
                                row_vals.append("TRUE" if val else "FALSE")
                            elif isinstance(val, (int, float)):
                                row_vals.append(str(val))
                            else:
                                clean_v = str(val).replace("'", "''")
                                row_vals.append(f"'{clean_v}'")
                        values_rows.append(f"({', '.join(row_vals)})")
                    
                    f_ins.write(f"INSERT INTO seduc.{table_name} ({cols_sql}) VALUES\n" + ",\n".join(values_rows) + ";\n")

        f_ins.write("\nCOMMIT;\n")

    # Salva schema_seduc.sql
    with open(schema_sql_path, "w", encoding="utf-8") as f_sch:
        f_sch.write("\n".join(schema_statements))

    # Salva load_data.sql
    with open(load_sql_path, "w", encoding="utf-8") as f_ld:
        f_ld.write("\n".join(load_statements))

    print(f"\n[PostgreSQL] Artefatos gerados com sucesso na pasta: {PG_DIR}")
    print(f"  -> Esquema DDL: {schema_sql_path}")
    print(f"  -> Carga via COPY: {load_sql_path}")
    print(f"  -> Carga via INSERTS: {inserts_sql_path}")
    print(f"  -> CSVs limpos: {PG_DATA_DIR}")


# ----------------------------------------------------------------------
# EXECUÇÃO PRINCIPAL DO PIPELINE
# ----------------------------------------------------------------------

def main():
    print("=" * 80)
    print("INICIANDO PIPELINE DE EXTRAÇÃO SEDUC: PDF -> EXCEL (ABAS) & POSTGRESQL")
    print("=" * 80)

    extracted_datasets = []
    summary_data = []

    # Cria novo Workbook do Excel
    wb = openpyxl.Workbook()
    # Remove sheet inicial padrao
    wb.remove(wb.active)

    for idx, cfg in enumerate(DATASET_CONFIG, 1):
        fname = cfg["filename"]
        fpath = os.path.join(PDF_DIR, fname)
        sheet_name = cfg["sheet_name"]
        table_name = cfg["table_name"]
        title = cfg["title"]

        print(f"\n[{idx}/9] Processando: {fname}...")

        if not os.path.exists(fpath):
            print(f"  [AVISO] Arquivo não encontrado: {fpath}")
            continue

        extractor = cfg["extractor"]
        df = extractor(fpath)
        print(f"  [OK] {len(df)} registros extraídos.")

        # Adiciona no Excel como aba individual
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        # Preenche os dados
        # Linha de cabeçalho
        ws.append(list(df.columns))
        # Linhas de dados
        for row in df.itertuples(index=False):
            ws.append(list(row))

        # Aplica estilos visuais profissionais
        style_worksheet(ws, df, title)

        extracted_datasets.append({
            "filename": fname,
            "sheet_name": sheet_name,
            "table_name": table_name,
            "title": title,
            "df": df
        })

        summary_data.append({
            "filename": fname,
            "sheet_name": sheet_name,
            "table_name": table_name,
            "title": title,
            "records": len(df)
        })

    # Cria a aba de Visão Geral / Painel Inicial
    print("\n[Excel] Criando aba de Visão Geral (00_Visao_Geral)...")
    create_overview_sheet(wb, summary_data)

    # Salva o arquivo Excel consolidado
    print(f"[Excel] Salvando planilha consolidada em: {OUTPUT_EXCEL}...")
    wb.save(OUTPUT_EXCEL)
    print("  [OK] Planilha Excel salva com sucesso!")

    # Gera artefatos para o PostgreSQL
    print("\n[PostgreSQL] Gerando scripts DDL, CSVs e INSERTS para PostgreSQL...")
    generate_postgres_artifacts(extracted_datasets)

    print("\n" + "=" * 80)
    print("PIPELINE CONCLUÍDO COM SUCESSO!")
    print(f"1. Planilha Excel: {OUTPUT_EXCEL}")
    print(f"2. Scripts e Dados PostgreSQL: {PG_DIR}")
    print("=" * 80)


if __name__ == "__main__":
    main()
